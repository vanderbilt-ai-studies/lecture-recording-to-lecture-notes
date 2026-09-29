#!/usr/bin/env python3
"""Extract structured text from PPTX slides using python-pptx."""

import argparse
import json
import sys

try:
    from pptx import Presentation
except ImportError:
    print("python-pptx not found. Installing...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
    from pptx import Presentation


def extract(pptx_path: str) -> list[dict]:
    prs = Presentation(pptx_path)
    slides_data = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
        slides_data.append({"slide_number": i + 1, "texts": texts})
    return slides_data


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract text from PPTX slides")
    parser.add_argument("pptx_path", help="Path to the PPTX file")
    parser.add_argument("--output", "-o", help="Output JSON file (default: stdout)")
    args = parser.parse_args()

    data = extract(args.pptx_path)
    output = json.dumps(data, indent=2)

    if args.output:
        with open(args.output, "w") as f:
            f.write(output)
        print(f"Extracted text from {len(data)} slides to {args.output}")
    else:
        print(output)
